import io
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "WS2.0"))
import app as application


def test_download_all_contacts_exports_saved_collection(monkeypatch, tmp_path):
    master_file = tmp_path / "master_email_list.csv"
    lock_file = tmp_path / "data.lock"
    pd.DataFrame([
        {
            "Paper Title": "Research paper",
            "Author Name": "Jane Smith",
            "Email ID": "jane@university.edu",
        }
    ]).to_csv(master_file, index=False)
    monkeypatch.setattr(application, "MASTER_EMAIL_FILE", str(master_file))
    monkeypatch.setattr(application, "DATA_LOCK_FILE", str(lock_file))

    response = application.app.test_client().get("/download/all")

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["Content-Disposition"].endswith('filename=medscout_all_contacts.xlsx')
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    worksheet = workbook.active
    assert list(worksheet.values) == [
        ("Paper Title", "Author Name", "Email ID"),
        ("Research paper", "Jane Smith", "jane@university.edu"),
    ]


def test_download_all_contacts_reports_empty_collection(monkeypatch, tmp_path):
    monkeypatch.setattr(application, "MASTER_EMAIL_FILE", str(tmp_path / "missing.csv"))
    monkeypatch.setattr(application, "DATA_LOCK_FILE", str(tmp_path / "data.lock"))

    response = application.app.test_client().get("/download/all")

    assert response.status_code == 404
    assert response.get_json()["error"] == "There are no saved contacts to export yet."


def test_download_current_results_creates_workbook_without_saving(monkeypatch, tmp_path):
    rows = [{
        "Paper Title": "Current paper",
        "Author Name": "John Doe",
        "Email ID": "john@college.edu",
    }]
    monkeypatch.setattr(application, "DOWNLOADS_DIR", str(tmp_path))
    response = application.app.test_client().post("/download/current", json={"data": rows})

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    assert list(workbook.active.values) == [
        ("Paper Title", "Author Name", "Email ID"),
        ("Current paper", "John Doe", "john@college.edu"),
    ]
    assert not list(tmp_path.rglob("*.xlsx"))
